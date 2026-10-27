from __future__ import annotations

import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Color, Font, PatternFill, Side

from pbi_client import InvoiceLine

PROMPT_PAY_RATE = 0.65

_DATE_FMT = "[$-10409]mm/dd/yy"
_MONEY_FMT = '[$-10409]"$"0.00;\\("$"0.00\\)'
_PROMPT_FMT = '"$"#,##0.00'

_THIN = Side(style="thin", color=Color(indexed=64))
_BOX = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_OPEN_RIGHT = Border(left=_THIN, top=_THIN, bottom=_THIN)
_OPEN_SIDES = Border(top=_THIN, bottom=_THIN)
_OPEN_LEFT = Border(right=_THIN, top=_THIN, bottom=_THIN)

_FACE = "Aptos Narrow"
_BASE = Font(name=_FACE, size=11, family=2, color=Color(theme=1))
_TITLE = Font(name=_FACE, size=20, family=2, color=Color(theme=1))
_HEAD = Font(name=_FACE, size=12, family=2, bold=True, color="FF666666")
_BODY = Font(name=_FACE, size=11, family=2, color="FF333333")
_BOLD = Font(name=_FACE, size=11, family=2, bold=True, color="FF333333")
_WHITE = PatternFill("solid", fgColor="FFFFFFFF", bgColor="FFFFFFFF")

_CENTER = Alignment(horizontal="center")
_HEAD_ALIGN = Alignment(horizontal="center", wrap_text=True, readingOrder=1)
_MID_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True, readingOrder=1)
_MID_RIGHT = Alignment(horizontal="right", vertical="center", wrap_text=True, readingOrder=1)
_MID_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True, readingOrder=1)

_WIDTHS = {
    "A": 20.5703125,
    "B": 19.28515625,
    "C": 24.140625,
    "E": 23.42578125,
    "F": 24.5703125,
    "G": 25.5703125,
    "H": 29.28515625,
}


def _put(ws, ref, value=None, *, font=_BASE, fill=None, border=None, align=None, fmt=None):
    cell = ws[ref]
    if value is not None:
        cell.value = value
    cell.font = font
    if fill is not None:
        cell.fill = fill
    if border is not None:
        cell.border = border
    if align is not None:
        cell.alignment = align
    if fmt is not None:
        cell.number_format = fmt


def _merge_row(ws, row: int) -> None:
    ws.merge_cells(f"A{row}:B{row}")
    ws.merge_cells(f"D{row}:E{row}")


def _trip_date(value: str):
    text = (value or "").strip()
    for pattern, width in (("%Y-%m-%d", 10), ("%m/%d/%Y", 10)):
        try:
            return datetime.strptime(text[:width], pattern)
        except ValueError:
            continue
    return text


def summary_title(invoice: str) -> str:
    return f"NORTHBAY CLAIMS SUMMARY {invoice[4:6]}.{invoice[6:8]}.{invoice[0:4]}"


def build_summary(invoice: str, lines: list[InvoiceLine]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.sheet_view.topLeftCell = "B1"
    for col, width in _WIDTHS.items():
        ws.column_dimensions[col].width = width
    ws.column_dimensions["A"].hidden = True

    ws.row_dimensions[1].height = 31.5
    ws.merge_cells("A1:H1")
    _put(ws, "A1", summary_title(invoice), font=_TITLE, border=_OPEN_RIGHT, align=_CENTER)
    for col in "BCDEFG":
        _put(ws, f"{col}1", font=_TITLE, border=_OPEN_SIDES, align=_CENTER)
    _put(ws, "H1", font=_TITLE, border=_OPEN_LEFT, align=_CENTER)

    head = {"font": _HEAD, "fill": _WHITE, "align": _HEAD_ALIGN}
    _merge_row(ws, 2)
    _put(ws, "A2", "DATE", border=_OPEN_RIGHT, **head)
    _put(ws, "B2", border=_OPEN_LEFT, **head)
    _put(ws, "C2", "RUN", border=_BOX, **head)
    _put(ws, "D2", "NAME", border=_OPEN_RIGHT, **head)
    _put(ws, "E2", border=_OPEN_LEFT, **head)
    _put(ws, "F2", "CHARGE", border=_BOX, **head)
    _put(ws, "G2", "BALANCE", border=_BOX, **head)
    _put(ws, "H2", "PROMPT PAY-30 DAYS ", border=_BOX, **head)

    body = {"font": _BODY, "fill": _WHITE}
    row = 3
    for line in lines:
        _merge_row(ws, row)
        _put(ws, f"A{row}", _trip_date(line.trip_date), border=_OPEN_RIGHT, align=_MID_CENTER, fmt=_DATE_FMT, **body)
        _put(ws, f"B{row}", border=_OPEN_LEFT, align=_MID_CENTER, fmt=_DATE_FMT, **body)
        _put(ws, f"C{row}", line.run, border=_BOX, align=_MID_CENTER, **body)
        _put(ws, f"D{row}", line.name, border=_OPEN_RIGHT, align=_MID_CENTER, **body)
        _put(ws, f"E{row}", border=_OPEN_LEFT, align=_MID_CENTER, **body)
        _put(ws, f"F{row}", line.charge, border=_BOX, align=_MID_CENTER, **body)
        _put(ws, f"G{row}", round(line.billed, 2), border=_BOX, align=_MID_RIGHT, fmt=_MONEY_FMT, **body)
        _put(ws, f"H{row}", f"=G{row}*{PROMPT_PAY_RATE}", border=_BOX, fmt=_PROMPT_FMT)
        row += 1

    total = {"font": _BOLD, "fill": _WHITE}
    _merge_row(ws, row)
    _put(ws, f"A{row}", "TOTAL", border=_OPEN_RIGHT, align=_MID_CENTER, **total)
    _put(ws, f"B{row}", border=_OPEN_LEFT, align=_MID_CENTER, **total)
    _put(ws, f"C{row}", "", border=_BOX, align=_MID_LEFT, **total)
    _put(ws, f"D{row}", "", border=_OPEN_RIGHT, align=_MID_LEFT, **total)
    _put(ws, f"E{row}", border=_OPEN_LEFT, align=_MID_LEFT, **total)
    _put(ws, f"F{row}", "", border=_BOX, align=_MID_CENTER, **total)
    _put(ws, f"G{row}", round(sum(line.billed for line in lines), 2), border=_BOX, align=_MID_RIGHT, fmt=_MONEY_FMT, **total)
    _put(ws, f"H{row}", f"=G{row}*{PROMPT_PAY_RATE}", border=_BOX, fmt=_PROMPT_FMT)
    row += 1

    _merge_row(ws, row)
    _put(ws, f"A{row}", border=_OPEN_RIGHT, align=_CENTER)
    _put(ws, f"B{row}", border=_OPEN_LEFT, align=_CENTER)
    _put(ws, f"C{row}", border=_BOX)
    _put(ws, f"D{row}", border=_OPEN_RIGHT, align=_CENTER)
    _put(ws, f"E{row}", border=_OPEN_LEFT, align=_CENTER)
    for col in "FGH":
        _put(ws, f"{col}{row}", border=_BOX)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
